import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Margins 2.5cm
for section in doc.sections:
    section.top_margin = Inches(0.98)
    section.bottom_margin = Inches(0.98)
    section.left_margin = Inches(0.98)
    section.right_margin = Inches(0.98)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_title(text, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
    return p

def add_body(text, bold_prefix="", italic=False, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    r.font.italic = italic
    return p

import os

def add_placeholder_box(figura_num, legenda, image_path=None):
    if image_path and os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(image_path, width=Inches(6.0))
    else:
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=350, bottom=350, left=200, right=200)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run("[ COLE O PRINT AQUI ]\n")
        r1.font.name = "Arial"
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)

        r2 = p.add_run("Clique nesta caixa e cole a imagem (Ctrl + V)")
        r2.font.name = "Arial"
        r2.font.size = Pt(9)
        r2.font.italic = True
        r2.font.color.rgb = RGBColor(0x94, 0xa3, 0xb8)

    p_legenda = doc.add_paragraph()
    p_legenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_legenda.paragraph_format.space_before = Pt(4)
    p_legenda.paragraph_format.space_after = Pt(14)
    r_leg = p_legenda.add_run(f"Figura {figura_num}: {legenda}")
    r_leg.font.name = "Times New Roman"
    r_leg.font.size = Pt(9.5)
    r_leg.font.italic = True

def add_ficha(descricao, planejada, realizada, situacao, evidencia, link):
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths = [Inches(1.8), Inches(4.7)]
    rows_data = [
        ("Descrição", descricao),
        ("Data planejada", planejada),
        ("Data realizada", realizada),
        ("Situação", situacao),
        ("Evidência", evidencia),
        ("Link", link),
    ]

    for i, (campo, valor) in enumerate(rows_data):
        c0 = table.cell(i, 0)
        c1 = table.cell(i, 1)
        c0.width = col_widths[0]
        c1.width = col_widths[1]
        set_cell_background(c0, "F8FAFC")
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(campo)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(9.5)
        r0.font.bold = True

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(valor)
        r1.font.name = "Times New Roman" if campo != "Link" else "Courier New"
        r1.font.size = Pt(9.5 if campo != "Link" else 8.5)
        if campo == "Situação":
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ══════════════════════════════════════════════════════
# CAPA
# ══════════════════════════════════════════════════════
p_inst = doc.add_paragraph()
p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_inst.paragraph_format.space_before = Pt(80)
p_inst.paragraph_format.space_after = Pt(2)
r_inst = p_inst.add_run("FATEC POMPEIA")
r_inst.font.name = "Times New Roman"
r_inst.font.size = Pt(13)
r_inst.font.bold = True

p_disc = doc.add_paragraph()
p_disc.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_disc.paragraph_format.space_after = Pt(120)
r_disc = p_disc.add_run("Introdução à Computação em Nuvem – ISW055 – 161_SIST. INTELIGENTES_N")
r_disc.font.name = "Times New Roman"
r_disc.font.size = Pt(10.5)

p_p1 = doc.add_paragraph()
p_p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_p1.paragraph_format.space_after = Pt(4)
r_p1 = p_p1.add_run("P1")
r_p1.font.name = "Times New Roman"
r_p1.font.size = Pt(28)
r_p1.font.bold = True

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_after = Pt(2)
r_sub = p_sub.add_run("Relatório bimestral de atividades")
r_sub.font.name = "Times New Roman"
r_sub.font.size = Pt(18)

p_av = doc.add_paragraph()
p_av.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_av.paragraph_format.space_after = Pt(130)
r_av = p_av.add_run("Avaliação individual – 2026.2")
r_av.font.name = "Times New Roman"
r_av.font.size = Pt(11)
r_av.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)

p_nome = doc.add_paragraph()
p_nome.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_nome.paragraph_format.space_after = Pt(100)
r_nome = p_nome.add_run("Lucas Tetsuya Hiratuca")
r_nome.font.name = "Times New Roman"
r_nome.font.size = Pt(14)
r_nome.font.bold = True

p_prof = doc.add_paragraph()
p_prof.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_prof.paragraph_format.space_after = Pt(2)
r_prof = p_prof.add_run("Prof. Allan Lincoln Rodrigues Siriani")
r_prof.font.name = "Times New Roman"
r_prof.font.size = Pt(10.5)

p_data = doc.add_paragraph()
p_data.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_data = p_data.add_run("Pompeia, 07/10/2026")
r_data.font.name = "Times New Roman"
r_data.font.size = Pt(10.5)

doc.add_page_break()

# ══════════════════════════════════════════════════════
# SUMÁRIO E SEÇÕES
# ══════════════════════════════════════════════════════
add_title("Sumário", size=16)
itens_sumario = [
    "1  Introdução",
    "2  Metodologia",
    "3  Quadro de entregas",
    "4  Atividades realizadas",
    "    4.1  Atividade 1 – Agenda telefônica em Flask",
    "    4.2  Atividade 2 – Catálogo de filmes – Tom Hanks",
    "    4.3  Atividade 3 – Desacoplando o login – microsserviço de autenticação",
    "    4.4  Atividade 4 – Controle de acesso por papel – RBAC",
    "    4.5  Atividade 5 – Logs e auditoria",
    "    4.6  Atividade 6 – Upload e perfil de usuário",
    "5  Considerações finais",
    "6  Declaração de autoria",
]
for item in itens_sumario:
    p_s = doc.add_paragraph()
    p_s.paragraph_format.space_after = Pt(2)
    r_s = p_s.add_run(item)
    r_s.font.name = "Times New Roman"
    r_s.font.size = Pt(10)

doc.add_page_break()

add_title("1  Introdução", size=14)
add_body("A disciplina ISW055 – Introdução à Computação em Nuvem, ministrada pelo Prof. Allan Lincoln Rodrigues Siriani na FATEC Pompeia no semestre letivo 2026.2, tem como propósito central proporcionar aos alunos a vivência prática e conceitual na construção, conteinerização e orquestração de aplicações modernas para ambientes em nuvem.")
add_body("O projeto condutor ao longo deste primeiro bimestre foi o desenvolvimento e a evolução do Catálogo de Filmes de Tom Hanks. Partindo de uma arquitetura inicial monolítica, a solução progrediu de forma incremental para uma infraestrutura desacoplada em microsserviços conteinerizados com Docker. Foram incorporados requisitos reais de engenharia de software em nuvem: autenticação segura via JSON Web Tokens (JWT), controle de acesso baseado em papéis (RBAC), auditoria assíncrona de eventos com Redis Streams e persistência de mídia em object storage compatível com S3 (Garage S3).")
add_body("Este relatório bimestral documenta minuciosamente as seis atividades propostas e entregues, detalhando as decisões de arquitetura adotadas, as dificuldades técnicas superadas, as evidências de entrega com carimbo temporal no Git e os resultados operacionais comprovados do sistema.")

add_title("2  Metodologia", size=14)
add_body("Todas as atividades foram desenvolvidas com controle de versão via Git em repositórios públicos no GitHub. As datas e os horários de entrega foram rigorosamente extraídos do histórico de commits (git log) e conferidos diretamente na interface web do GitHub, garantindo rastreabilidade e imutabilidade das comprovações.")
add_body("O ambiente de execução e teste foi estruturado com Docker Desktop e Docker Compose, padronizando os contêineres entre os ambientes de desenvolvimento e produção. O deploy final em nuvem é gerenciado via Portainer no servidor disponibilizado pelo orientador (lucas-hiratuca-isw055.lapps.studio), exposto por túnel seguro Cloudflare. As imagens oficiais foram publicadas no Docker Hub no namespace lucashiratuca/.")

add_title("3  Quadro de entregas", size=14)
quadro_data = [
    ("Nº", "Atividade", "Descrição", "Data planejada", "Data realizada", "Situação"),
    ("1", "Agenda telefônica em Flask", "Nivelamento em sala: sistema monolítico Flask + Jinja com persistência em MariaDB.", "07/08/2026", "07/08/2026 20:03", "entregue"),
    ("2", "Catálogo de filmes – Tom Hanks", "Consumo da API TMDB, persistência em MariaDB e segregação por usuário.", "20/08/2026", "20/08/2026 20:54", "entregue"),
    ("3", "Desacoplando o login", "Login, cadastro e esqueci-minha-senha num serviço à parte na rede interna do Docker.", "28/08/2026", "27/08/2026 13:00", "entregue"),
    ("4", "Controle de acesso – RBAC", "O campo role passa a decidir permissões reais no backend (403 para usuário comum).", "04/09/2026", "01/09/2026 00:48", "entregue"),
    ("5", "Logs e auditoria", "Novo log-service com Redis registrando login, ações sensíveis e tentativas negadas.", "25/09/2026", "08/09/2026 12:57", "entregue"),
    ("6", "Upload e perfil de usuário", "Página de perfil com avatar no Garage S3; só a referência fica no banco relacional.", "02/10/2026", "26/09/2026 14:43", "entregue"),
]
t_q = doc.add_table(rows=len(quadro_data), cols=6)
t_q.alignment = WD_TABLE_ALIGNMENT.CENTER
for r_idx, row in enumerate(quadro_data):
    for c_idx, val in enumerate(row):
        c = t_q.cell(r_idx, c_idx)
        if r_idx == 0:
            set_cell_background(c, "1E293B")
        elif r_idx % 2 == 0:
            set_cell_background(c, "F8FAFC")
        set_cell_margins(c, top=60, bottom=60, left=80, right=80)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (c_idx == 0 or c_idx >= 3) else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(val)
        run.font.name = "Times New Roman"
        run.font.size = Pt(8.5)
        if r_idx == 0:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        elif c_idx == 5:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x16, 0x65, 0x34)

doc.add_page_break()

# ══════════════════════════════════════════════════════
# ATIVIDADES REALIZADAS
# ══════════════════════════════════════════════════════
add_title("4  Atividades realizadas", size=14)

# 4.1
add_title("4.1  Atividade 1 – Agenda telefônica em Flask", size=12)
add_ficha(
    "Nivelamento em sala: sistema monolítico Flask + Jinja com persistência em banco de dados MariaDB na nuvem.",
    "07/08/2026", "07/08/2026 20:03", "entregue",
    "GitHub – commits 28230dc e 738ce73 + README com menção ao professor + print do sistema",
    "https://github.com/LucasHiratuca/aula_01_cloud"
)
add_body("Durante a aula inicial de nivelamento, foi desenvolvida uma aplicação web em Python utilizando o framework Flask e o motor de templates Jinja2. A aplicação implementa um CRUD de contatos telefônicos conectado a uma instância remota do MariaDB hospedada no IP 35.226.64.52.\n\nComo decisão técnica, a camada de acesso a dados foi isolada utilizando a biblioteca PyMySQL, configurando codificação utf8mb4 e tratamento de erros de conexão. Posteriormente, foi aplicado saneamento de credenciais via variáveis de ambiente com arquivo .gitignore (commit fa5f58d).", bold_prefix="O que foi feito. ")
add_placeholder_box(1, "Atividade 1 – evidência da entrega (link, data e hora visíveis)")
add_placeholder_box(2, "Atividade 1 – resultado (o sistema/mapa/artigo funcionando)")
add_body("A principal dificuldade inicial consistiu na latência e na autenticação segura contra a base de dados remota do laboratório. O problema foi sanado com o correto ajuste da porta MySQL e a parametrização das variáveis no ambiente local.", bold_prefix="Dificuldades e como foram resolvidas. ")

doc.add_page_break()

# 4.2
add_title("4.2  Atividade 2 – Catálogo de filmes – Tom Hanks", size=12)
add_ficha(
    "Consumo da API TMDB, persistência em MariaDB e segregação por usuário com sistema de favoritos.",
    "20/08/2026", "20/08/2026 20:54", "entregue",
    "GitHub – commit 9a57391 + README com menção ao professor + print do catálogo",
    "https://github.com/LucasHiratuca/Entrega-da-Atividade/commit/9a57391"
)
add_body("Criação da aplicação principal do projeto utilizando Node.js, Express e EJS. A aplicação consome a API pública do The Movie Database (TMDB), filtrando e exibindo a filmografia completa do ator Tom Hanks com posters, títulos originais, sinopses e notas de avaliação.\n\nFoi construído o sistema de autenticação local com express-session e senhas criptografadas com bcryptjs. Adicionalmente, foi implementada a funcionalidade de filmes favoritos vinculada ao ID do usuário autenticado no MariaDB, segregando a lista de favoritos por sessão ativa.", bold_prefix="O que foi feito. ")
add_placeholder_box(3, "Atividade 2 – evidência da entrega (link, data e hora visíveis)")
add_placeholder_box(4, "Atividade 2 – resultado (o sistema/mapa/artigo funcionando)")
add_body("A API do TMDB utiliza paginação de resultados. Para garantir que toda a filmografia do ator fosse renderizada sem cortes, foi implementada uma rotina assíncrona que itera sobre as páginas da API e consolida os dados em memória antes da renderização EJS.", bold_prefix="Dificuldades e como foram resolvidas. ")

doc.add_page_break()

# 4.3
add_title("4.3  Atividade 3 – Desacoplando o login – microsserviço de autenticação", size=12)
add_ficha(
    "Login, cadastro e esqueci-minha-senha num serviço à parte na rede interna do Docker.",
    "28/08/2026", "27/08/2026 13:00", "entregue",
    "GitHub – commit 320028a + docker-compose.yml + print do login funcionando",
    "https://github.com/LucasHiratuca/Entrega-da-Atividade/commit/320028a"
)
add_body("A arquitetura do sistema foi refatorada sob o padrão Backend-for-Frontend (BFF). O serviço de catálogo (catalogo-service) foi mantido como único ponto de entrada público (porta 8217), enquanto toda a lógica de identidade, login, cadastro de novos usuários e redefinição de senha foi isolada no novo auth-service (porta interna 4000).\n\nO auth-service não expõe portas ao host, comunicando-se exclusivamente pela rede interna Docker app_network. Ele emite tokens JWT criptograficamente assinados contendo a identidade do usuário. O fluxo de recuperação de senha foi integrado ao serviço SMTP do Mailtrap para despacho de e-mails com tokens descartáveis com expiração de 30 minutos.", bold_prefix="O que foi feito. ")
add_placeholder_box(5, "Atividade 3 – evidência da entrega (link, data e hora visíveis)")
add_placeholder_box(6, "Atividade 3 – resultado (o sistema/mapa/artigo funcionando)")
add_body("Ajustar a comunicação segura entre contêineres na mesma rede virtual do Docker Compose sem depender de IPs estáticos. A resolução foi alcançada utilizando os nomes dos serviços como hostnames DNS internos do Docker (http://auth-service:4000).", bold_prefix="Dificuldades e como foram resolvidas. ")

doc.add_page_break()

# 4.4
add_title("4.4  Atividade 4 – Controle de acesso por papel – RBAC", size=12)
add_ficha(
    "O campo role passa a decidir permissões reais no backend (403 para usuário comum).",
    "04/09/2026", "01/09/2026 00:48", "entregue",
    "GitHub – commits 6636f4d e a8e31ef + print do 403 e da ação de admin",
    "https://github.com/LucasHiratuca/Entrega-da-Atividade/commit/6636f4d"
)
add_body("Implementou-se o modelo de autorização Role-Based Access Control (RBAC) adotando o Padrão B (claims embutidas diretamente no payload do JWT). O campo role (usuario ou admin) trafega assinado no token, permitindo verificação instantânea no catalogo-service sem overhead de consultas repetidas ao banco de dados.\n\nForam criados os middlewares requireLogin e requireAdmin. Usuários autenticados com perfil comum têm acesso restrito ao catálogo, favoritos e comentários próprios. Tentativas de acessar rotas de moderação ou painel de gestão de usuários (/admin/usuarios) são barradas com status HTTP 403 Forbidden diretamente na camada do servidor.", bold_prefix="O que foi feito. ")
add_placeholder_box(7, "Atividade 4 – evidência da entrega (link, data e hora visíveis)")
add_placeholder_box(8, "Atividade 4 – resultado (painel de administração de usuários e RBAC em funcionamento)", "outputs/admin_output.png")
add_body("Garantir que tipos de dados inteiros para userId não causassem falhas na comparação estrita (===) no JavaScript. Foi adotado parsing consistente com parseInt() em todos os middlewares e controladores de permissão.", bold_prefix="Dificuldades e como foram resolvidas. ")

doc.add_page_break()

# 4.5
add_title("4.5  Atividade 5 – Logs e auditoria", size=12)
add_ficha(
    "Novo log-service com Redis registrando login, ações sensíveis e tentativas negadas.",
    "25/09/2026", "08/09/2026 12:57", "entregue",
    "GitHub – commits e6c6682 e 8bc9faa + print da consulta de logs pelo admin",
    "https://github.com/LucasHiratuca/Entrega-da-Atividade/commit/e6c6682"
)
add_body("Foi introduzido um terceiro microsserviço dedicado, o log-service (porta 5000), acoplado a uma instância do Redis 7. O serviço implementa ingestão assíncrona de eventos via HTTP POST, persistindo as entradas de auditoria na estrutura de alta performance Redis Streams.\n\nO catalogo-service consome o utilitário de auditoria registrando eventos críticos: logins com sucesso, encerramentos de sessão, falhas de autenticação, exclusão de comentários e edições de perfil. A interface /admin/logs foi criada com restrição exclusiva a administradores para visualização e filtragem cronológica reversa dos registros em tempo real.", bold_prefix="O que foi feito. ")
add_placeholder_box(9, "Atividade 5 – evidência da entrega (link, data e hora visíveis)")
add_placeholder_box(10, "Atividade 5 – resultado (tela de consulta de logs de auditoria via Redis Streams)", "outputs/logs_output.png")
add_body("Garantir a sobrevivência dos logs a reinicializações de contêineres. Foi configurado um volume nomeado persistente para o contêiner do Redis no docker-compose.yml, além de timeout não-bloqueante no envio de logs para evitar impacto na experiência do usuário final caso o Redis ficasse temporariamente indisponível.", bold_prefix="Dificuldades e como foram resolvidas. ")

doc.add_page_break()

# 4.6
add_title("4.6  Atividade 6 – Upload e perfil de usuário", size=12)
add_ficha(
    "Página de perfil com avatar no Garage S3; só a referência fica no banco relacional.",
    "02/10/2026", "26/09/2026 14:43", "entregue",
    "GitHub – commits 4219f05, 27856a8 e 8ca3fb6 + print do perfil com foto + SECURITY_AUDIT.md",
    "https://github.com/LucasHiratuca/Entrega-da-Atividade/commit/27856a8"
)
add_body("Construção da página de perfil com suporte a upload de foto de avatar persistida em object storage. Foi integrado o Garage S3 (dxflrs/garage:v1.0.1), um engine leve e distribuído escrito em Rust compatível com AWS S3 API (Signature V4). A tabela perfis no MariaDB armazena apenas a biografia e a chave textual do objeto (foto_key), desacoplando o binário do banco relacional.\n\nA entrega da imagem foi arquitetada via streaming seguro pelo backend (/perfil/:userId/foto): o catalogo-service obtém o stream do Garage S3 via rede interna e o devolve ao navegador com cabeçalhos HTTP adequados (Content-Type validado e Cache-Control). O upload exige validação de MIME type (JPEG/PNG/WebP/GIF), limite de 5MB e geração de nomes via randomUUID(). Foi realizada ainda auditoria OWASP Top 10 com 100% de aprovação e remediação do CWE-640 (Password Reset Poisoning).", bold_prefix="O que foi feito. ")
add_placeholder_box(11, "Atividade 6 – evidência da entrega (link, data e hora visíveis)")
add_placeholder_box(12, "Atividade 6 – resultado (perfil do usuário exibindo foto de avatar carregada do Garage S3 e bio)", "outputs/foto_perfil.png")
add_body("Em outubro de 2025, a imagem oficial do MinIO foi descontinuada do Docker Hub, motivando a migração definitiva para o Garage S3. Além disso, URLs pré-assinadas geravam links com o host interno do contêiner (garage:3900), inacessível aos navegadores dos clientes. A substituição por streaming direto no backend resolveu a compatibilidade universalmente em ambientes locais e Cloudflare.", bold_prefix="Dificuldades e como foram resolvidas. ")

doc.add_page_break()

# ══════════════════════════════════════════════════════
# 5 CONSIDERAÇÕES FINAIS & 6 AUTORIA
# ══════════════════════════════════════════════
add_title("5  Considerações finais", size=14)
add_body("O percurso de desenvolvimento percorrido ao longo do primeiro bimestre de ISW055 permitiu consolidar na prática os fundamentos de arquiteturas em nuvem modernas. A transição de um script inicial monolítico em Python para um ecossistema completo composto por microsserviços conteinerizados em Node.js, mensageria com Redis e storage compatível com S3 proporcionou um aprendizado expressivo sobre segregação de responsabilidades e resiliência.")
add_body("A maior complexidade encontrada ocorreu na Atividade 6, onde a descontinuidade do MinIO no Docker Hub exigiu adaptabilidade imediata para selecionar, configurar e implantar o Garage S3 em Rust, ajustando o layout de cluster e o roteamento interno de rede. Essa experiência evidenciou o valor de adotar padrões abertos de mercado (como a S3 API), permitindo trocar a engine de armazenamento sem alterar as rotas da aplicação.")
add_body("Como aprendizado contínuo, destaca-se a relevância de conceber sistemas com foco em segurança desde a camada de design (Security by Design). A auditoria baseada no OWASP Top 10 permitiu antecipar vetores de ataque como enumeração de credenciais e envenenamento de links. Para o segundo bimestre, pretende-se aprofundar na instrumentação de métricas com Prometheus e automação de pipelines de entrega contínua (CI/CD).")

add_title("6  Declaração de autoria", size=14)
add_body("Declaro que este relatório foi elaborado por mim, individualmente, e que as evidências apresentadas correspondem a entregas de minha autoria, verificáveis nos links informados. Nas atividades realizadas em grupo, o conteúdo aqui descrito refere-se à minha participação.")

p_esp = doc.add_paragraph()
p_esp.paragraph_format.space_before = Pt(60)

t_ass = doc.add_table(rows=1, cols=2)
t_ass.alignment = WD_TABLE_ALIGNMENT.CENTER
c_l = t_ass.cell(0, 0)
c_r = t_ass.cell(0, 1)
c_l.width = Inches(3.5)
c_r.width = Inches(3.0)

p_l = c_l.paragraphs[0]
p_l.add_run("_________________________________________\n").font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
r_aluno = p_l.add_run("Lucas Tetsuya Hiratuca\n")
r_aluno.font.bold = True
p_l.add_run("Aluno(a)")

p_r = c_r.paragraphs[0]
p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_r.add_run("\n\nPompeia, 07/10/2026")

doc.save("P1_Relatorio_LucasHiratuca.docx")
print("DOCX gerado com sucesso!")
